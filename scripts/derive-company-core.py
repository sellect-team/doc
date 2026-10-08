# -*- coding: utf-8 -*-
"""회사소개서 본편 -> 연말정산(bibidi) 제외판 파생.

본편(company-intro-2026.pptx)이 바뀌면 이 스크립트를 다시 돌려 제외판을 새로 뽑는다.
생성물: docs/company/_source/company-intro-2026-core.pptx, company-intro-2026-core.titles.txt
이어서 convert-pptx.ps1 로 변환한다 (아래 사용법 참고).

  python scripts/derive-company-core.py
  .\\scripts\\convert-pptx.ps1 -Source "docs\\company\\_source\\company-intro-2026-core.pptx" -Category company `
      -Name company-intro-2026-core -DocTitle "세일링스톤 회사소개서 (연말정산 제외)" `
      -Description "회사 개요와 핵심 역량, aisight·Apps 제품 라인업, 구축 사례와 도입 절차 - 연말정산 솔루션을 뺀 판" -Version "1.0"
"""
import io, re, sys, shutil, pathlib
from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / 'docs/company/_source/company-intro-2026.pptx'
DST = ROOT / 'docs/company/_source/company-intro-2026-core.pptx'
TITLES_SRC = ROOT / 'docs/company/company-intro-2026.titles.txt'
TITLES_DST = ROOT / 'docs/company/company-intro-2026-core.titles.txt'

BAD = ('bibidi', '연말정산', '더홍시')

shutil.copyfile(SRC, DST)
p = Presentation(DST)

# 1) 제목 파일 기준으로 뺄 장을 정한다 (bibidi 3장)
titles = [l for l in io.open(TITLES_SRC, encoding='utf-8').read().split('\n') if l.strip()]
assert len(titles) == len(p.slides), f'제목 {len(titles)} != 슬라이드 {len(p.slides)}'
drop = [i for i, t in enumerate(titles) if 'bibidi' in t]
assert len(drop) == 3, drop
print('제외 장:', [f'{i+1}p {titles[i]}' for i in drop])

# 2) 남는 장의 하드코딩 쪽번호("- N -" 텍스트 상자)를 새 번호로
new_no = {}
k = 0
for i in range(len(p.slides)):
    if i in drop: continue
    k += 1; new_no[i] = k
for i, s in enumerate(p.slides):
    if i in drop: continue
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r'\s*-\s*\d+\s*-\s*', sh.text_frame.text):
            r = sh.text_frame.paragraphs[0].runs
            r[0].text = f'- {new_no[i]} -'
            for x in r[1:]: x.text = ''

# 3) 본문 언급 제거: 2p 목차 부제, 4p 사업 분야 항목, 맺음 SOLUTION 라인
def set_run_text(sh, old, new):
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            if old in r.text:
                r.text = r.text.replace(old, new); return True
    return False

hits = 0
for i, s in enumerate(p.slides):
    if i in drop: continue
    for sh in s.shapes:
        if not sh.has_text_frame: continue
        t = sh.text_frame.text
        if not any(b in t for b in BAD): continue
        if set_run_text(sh, ' · bibidi 연말정산', ''): hits += 1; continue          # 2p 목차, 맺음
        if set_run_text(sh, ' · bibidi 연말정산 · ', ' · '): hits += 1; continue
        # 4p 목록: 해당 문단 통째로 제거
        for para in list(sh.text_frame.paragraphs):
            if any(b in para.text for b in BAD):
                para._p.getparent().remove(para._p); hits += 1
print('본문 언급 정리:', hits, '곳')

# 4) 장 삭제 (관계까지 끊어 슬라이드 파트가 패키지에서 빠지게)
lst = p.slides._sldIdLst
for i in sorted(drop, reverse=True):
    sld = list(lst)[i]
    p.part.drop_rel(sld.rId)
    lst.remove(sld)
p.save(DST)

# 5) 남은 언급 검사
p2 = Presentation(DST)
left = [(i + 1, sh.text_frame.text[:50]) for i, s in enumerate(p2.slides) for sh in s.shapes
        if sh.has_text_frame and any(b in sh.text_frame.text for b in BAD)]
bad_no = [(i + 1, sh.text_frame.text) for i, s in enumerate(p2.slides) for sh in s.shapes
          if sh.has_text_frame and (m := re.fullmatch(r'\s*-\s*(\d+)\s*-\s*', sh.text_frame.text)) and int(m.group(1)) != i + 1]
print(f'결과 {len(p2.slides)}장 · 남은 언급 {left or "없음"} · 쪽번호 불일치 {bad_no or "없음"}')
io.open(TITLES_DST, 'w', encoding='utf-8').write('\n'.join(t for i, t in enumerate(titles) if i not in drop) + '\n')
sys.exit(1 if left or bad_no else 0)
